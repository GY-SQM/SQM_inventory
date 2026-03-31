"""
SQM v5.9.5 — Allocation 출고 예약 다이얼로그
=============================================

엑셀 업로드 → 파싱 미리보기 → 예약(RESERVED) 실행 → 현황 조회
"""
from engine_modules.constants import SAMPLE_MT_THRESHOLD, STATUS_RESERVED
import logging
import threading
import time
from engine_modules.constants import CUSTOMER_PRESETS  # v6.7.3
import tkinter as tk
from collections import Counter
from tkinter import BOTH, END, LEFT, RIGHT, VERTICAL, X, Y, filedialog, ttk

from ..utils.ui_constants import (
    CustomMessageBox,
    ThemeColors,
    setup_dialog_geometry_persistence,
)

try:
    from ..utils.gui_bootstrap import ScrolledFrame as _ScrolledFrame
except ImportError:
    _ScrolledFrame = None

logger = logging.getLogger(__name__)

# SAMPLE_MT_THRESHOLD는 engine_modules.constants에서 import (중앙 관리)


def _allocation_tonbag_sample_counts(rows: list) -> tuple:
    """Allocation 행에서 500kg 톤백 개수와 샘플(1kg) 개수 계산. (tonbag_500, sample_count)"""
    tonbag_500 = 0
    sample_count = 0
    for r in rows:
        qty = 0.0
        if hasattr(r, 'get'):
            qty = float(r.get('qty_mt') or 0)
        else:
            qty = float(getattr(r, 'qty_mt', 0) or 0)
        if qty >= SAMPLE_MT_THRESHOLD:
            tonbag_500 += int(round(qty * 1000 / 500))  # MT → 500kg 단위
        else:
            sample_count += 1
    return tonbag_500, sample_count


ALLOC_PREVIEW_COLUMNS = [
    ("lot_no",        "LOT NO",        110, "center"),
    ("sap_no",        "SAP NO",        100, "center"),
    ("product",       "PRODUCT",       140, "center"),
    ("qty_mt",        "QTY (MT)",       80, "center"),
    ("sold_to",       "CUSTOMER",      130, "center"),
    ("sale_ref",      "SALE REF",      120, "center"),
    ("outbound_date", "OUTBOUND DATE", 100, "center"),
    ("warehouse",     "WH",             60, "center"),
    ("status",        "STATUS",         80, "center"),
]


class AllocationDialog:
    """Allocation 엑셀 → 미리보기 → 예약/실행/취소"""

    def __init__(self, app, engine):
        self.app = app
        self.engine = engine
        self.root = getattr(app, 'root', None)
        self.dialog = None
        self.parsed_rows = []
        self.source_file = ""
        self._cell_editor = None
        self._cell_editor_ctx = None
        self._tree_ctx_menu = None
        self._editable_cols = {'lot_no', 'sap_no', 'product', 'qty_mt', 'sold_to', 'sale_ref', 'outbound_date', 'warehouse'}
        default_lot_mode = False
        try:
            if hasattr(self.engine, '_get_allocation_reservation_mode'):
                default_lot_mode = self.engine._get_allocation_reservation_mode() == "lot"
        except Exception:
            default_lot_mode = False
        self._lot_mode_var = tk.BooleanVar(value=default_lot_mode)

    def show(self, initial_file: str = None):
        """다이얼로그 표시. initial_file이 있으면 해당 파일 로드 후 파싱"""
        self.dialog = tk.Toplevel(self.root)
        self.dialog.title("📋 판매 배정 출고 예약")
        setup_dialog_geometry_persistence(self.dialog, "allocation_dialog", self.root, "large")
        self.dialog.transient(self.root)
        self.dialog.grab_set()
        _is_dark = ThemeColors.is_dark_theme(getattr(self.app, 'current_theme', 'flatly'))
        _bg = ThemeColors.get('bg_card', _is_dark)
        self.dialog.configure(bg=_bg)
        self._create_widgets()
        self.dialog.update_idletasks()
        if initial_file:
            self._file_var.set(initial_file)
            self.source_file = initial_file
            self.dialog.after(100, self._parse_file)

    def show_with_data(self, rows: list):
        """붙여넣기 등으로 받은 데이터(dict 리스트)로 다이얼로그 표시. 파일 없이 미리보기 → 예약."""
        self.parsed_rows = rows
        self.source_file = "(붙여넣기)"
        self.dialog = tk.Toplevel(self.root)
        self.dialog.title("📋 판매 배정 출고 예약")
        setup_dialog_geometry_persistence(self.dialog, "allocation_dialog", self.root, "large")
        self.dialog.transient(self.root)
        self.dialog.grab_set()
        _is_dark = ThemeColors.is_dark_theme(getattr(self.app, 'current_theme', 'flatly'))
        _bg = ThemeColors.get('bg_card', _is_dark)
        self.dialog.configure(bg=_bg)
        self._create_widgets()
        self.dialog.update_idletasks()
        self._file_var.set("(붙여넣기 데이터)")
        self._fill_tree_from_parsed_rows()
        total_mt = sum(float(r.get('qty_mt') or 0) for r in rows)
        tb500, samp = _allocation_tonbag_sample_counts(rows)
        self._summary_var.set(
            f"고객: (붙여넣기) | 총 {len(rows)}행 | 총량: {total_mt:.4f} MT | 500kg {tb500}개, 샘플 {samp}개"
        )
        if self.parsed_rows:
            self.btn_reserve.config(state='normal')

    def _create_widgets(self):
        top = ttk.Frame(self.dialog, padding=8)
        top.pack(fill=X)

        ttk.Label(top, text="판매 배정 Excel 파일:").pack(side=LEFT, padx=(0, 5))
        self._file_var = tk.StringVar()
        ttk.Entry(top, textvariable=self._file_var, width=60, state='readonly').pack(side=LEFT, fill=X, expand=True, padx=(0, 5))
        ttk.Button(top, text="📂 파일 선택", command=self._select_file).pack(side=LEFT, padx=(0, 5))

        # 데이터가 많을 때 창 전체를 스크롤할 수 있도록 본문을 ScrolledFrame으로 감쌈
        if _ScrolledFrame is not None:
            body_container = _ScrolledFrame(self.dialog, autohide=True)
            body_container.pack(fill=BOTH, expand=True)
        else:
            body_container = ttk.Frame(self.dialog)
            body_container.pack(fill=BOTH, expand=True)

        tree_frame = ttk.Frame(body_container, padding=8)
        tree_frame.pack(fill=BOTH, expand=True)

        cols = [c[0] for c in ALLOC_PREVIEW_COLUMNS]
        self.tree = ttk.Treeview(tree_frame, columns=cols, show='headings', height=18)
        for col_id, header, width, anchor in ALLOC_PREVIEW_COLUMNS:
            self.tree.heading(col_id, text=header)
            self.tree.column(col_id, width=width, anchor=anchor, minwidth=40)

        vsb = tk.Scrollbar(tree_frame, orient=VERTICAL, command=self.tree.yview)
        hsb = tk.Scrollbar(tree_frame, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        # AllocationDialog는 자체 편집 로직 사용(전역 Editable Treeview와 중복 방지)
        self.tree._disable_global_editable = True
        self.tree.pack(side=LEFT, fill=BOTH, expand=True)
        vsb.pack(side=RIGHT, fill=Y)
        hsb.pack(side=tk.BOTTOM, fill=tk.X)
        self._setup_editable_tree_bindings()
        try:
            from ..utils.tree_enhancements import TreeviewTotalFooter
            self._alloc_total_footer = TreeviewTotalFooter(
                tree_frame, self.tree, ['qty_mt'],
                column_display_names={'qty_mt': 'QTY (MT)'},
                column_formats={'qty_mt': ',.4f'},
            )
            self._alloc_total_footer.pack(fill=tk.X, pady=(2, 0))
        except (ImportError, Exception):
            self._alloc_total_footer = None

        self._summary_var = tk.StringVar(value="파일을 선택하세요")
        ttk.Label(body_container, textvariable=self._summary_var, padding=5).pack(fill=X)

        btn_frame = ttk.Frame(body_container, padding=8)
        btn_frame.pack(fill=X)

        self.btn_reserve = ttk.Button(btn_frame, text="✅ 예약 실행 (RESERVED)", command=self._on_reserve, state='disabled')
        self.btn_reserve.pack(side=LEFT, padx=5)

        self.btn_status = ttk.Button(btn_frame, text="📊 예약 현황", command=self._show_reservation_status)
        self.btn_status.pack(side=LEFT, padx=5)

        self.btn_cancel_res = ttk.Button(btn_frame, text="❌ 예약 취소", command=self._on_cancel_reservation, state='disabled')
        self.btn_cancel_res.pack(side=LEFT, padx=5)

        self.btn_reset_lot = ttk.Button(
            btn_frame, text="🧹 LOT 예약 초기화", command=self._on_reset_reservation_for_lots, state='disabled'
        )
        self.btn_reset_lot.pack(side=LEFT, padx=5)

        self.btn_execute = ttk.Button(btn_frame, text="📦 출고 실행 (PICKED)", command=self._on_execute, state='disabled')
        self.btn_execute.pack(side=LEFT, padx=5)

        self.btn_confirm = ttk.Button(btn_frame, text="🔒 출고 확정 (SOLD)", command=self._on_confirm, state='disabled')
        self.btn_confirm.pack(side=LEFT, padx=5)

        ttk.Checkbutton(
            btn_frame,
            text="LOT 단위 예약 모드(톤백 ID 미지정)",
            variable=self._lot_mode_var
        ).pack(side=LEFT, padx=8)

        ttk.Button(btn_frame, text="닫기", command=self.dialog.destroy).pack(side=RIGHT, padx=5)

    def _setup_editable_tree_bindings(self):
        """Allocation 미리보기 Treeview를 엑셀 유사 편집 모드로 확장."""
        self.tree.bind("<Double-1>", self._on_tree_double_click)
        self.tree.bind("<Delete>", self._delete_selected_rows)
        self.tree.bind("<Control-c>", self._copy_selected_rows)
        self.tree.bind("<Control-x>", self._cut_selected_rows)
        self.tree.bind("<Control-v>", self._paste_rows_from_clipboard)
        self.tree.bind("<Control-a>", self._select_all_rows)
        self.tree.bind("<Button-3>", self._open_tree_context_menu)

        self._tree_ctx_menu = tk.Menu(self.tree, tearoff=0)
        self._tree_ctx_menu.add_command(label="행 삭제", command=self._delete_selected_rows)
        self._tree_ctx_menu.add_command(label="복사", command=self._copy_selected_rows)
        self._tree_ctx_menu.add_command(label="잘라내기", command=self._cut_selected_rows)
        self._tree_ctx_menu.add_separator()
        self._tree_ctx_menu.add_command(label="전체 초기화", command=self._clear_all_rows)

    def _open_tree_context_menu(self, event):
        row_id = self.tree.identify_row(event.y)
        if row_id:
            self.tree.selection_set((row_id,))
        try:
            self._tree_ctx_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self._tree_ctx_menu.grab_release()

    def _on_tree_double_click(self, event):
        region = self.tree.identify("region", event.x, event.y)
        if region != "cell":
            return
        row_id = self.tree.identify_row(event.y)
        col_token = self.tree.identify_column(event.x)  # #1, #2...
        if not row_id or not col_token:
            return
        col_idx = int(str(col_token).replace("#", "")) - 1
        cols = list(self.tree["columns"])
        if col_idx < 0 or col_idx >= len(cols):
            return
        col_name = cols[col_idx]
        if col_name not in self._editable_cols:
            return

        bbox = self.tree.bbox(row_id, col_token)
        if not bbox:
            return
        x, y, w, h = bbox
        old_val = self.tree.set(row_id, col_name)

        self._cancel_cell_editor()
        self._cell_editor = ttk.Entry(self.tree)
        self._cell_editor.place(x=x, y=y, width=w, height=h)
        self._cell_editor.insert(0, old_val)
        self._cell_editor.focus_set()
        self._cell_editor.selection_range(0, tk.END)
        self._cell_editor_ctx = (row_id, col_name)
        self._cell_editor.bind("<Return>", self._commit_cell_editor)
        self._cell_editor.bind("<Escape>", self._cancel_cell_editor)
        self._cell_editor.bind("<FocusOut>", self._commit_cell_editor)

    def _commit_cell_editor(self, _event=None):
        if not self._cell_editor or not self._cell_editor_ctx:
            return
        row_id, col_name = self._cell_editor_ctx
        new_val = self._cell_editor.get().strip()
        if col_name == 'qty_mt':
            try:
                num = float(new_val.replace(",", "")) if new_val else 0.0
                new_val = f"{num:.4f}"
            except (ValueError, TypeError):
                new_val = "0.0000"
        self.tree.set(row_id, col_name, new_val)
        self._cancel_cell_editor()
        self._after_tree_data_changed()

    def _cancel_cell_editor(self, _event=None):
        if self._cell_editor:
            try:
                self._cell_editor.destroy()
            except Exception as _ce:
                logging.getLogger(__name__).debug(f"[Allocation] 셀 에디터 정리 실패: {_ce}")
        self._cell_editor = None
        self._cell_editor_ctx = None

    def _selected_iids(self):
        return list(self.tree.selection() or ())

    def _copy_selected_rows(self, _event=None):
        iids = self._selected_iids()
        if not iids:
            return "break"
        cols = list(self.tree["columns"])
        lines = []
        for iid in iids:
            vals = self.tree.item(iid, "values")
            line = "\t".join(str(vals[i]) if i < len(vals) else "" for i in range(len(cols)))
            lines.append(line)
        txt = "\n".join(lines)
        self.tree.clipboard_clear()
        self.tree.clipboard_append(txt)
        return "break"

    def _cut_selected_rows(self, _event=None):
        self._copy_selected_rows()
        self._delete_selected_rows()
        return "break"

    def _delete_selected_rows(self, _event=None):
        iids = self._selected_iids()
        for iid in iids:
            self.tree.delete(iid)
        self._after_tree_data_changed()
        return "break"

    def _select_all_rows(self, _event=None):
        all_iids = self.tree.get_children("")
        if all_iids:
            self.tree.selection_set(all_iids)
        return "break"

    def _clear_all_rows(self):
        for iid in self.tree.get_children(""):
            self.tree.delete(iid)
        self._after_tree_data_changed()

    def _paste_rows_from_clipboard(self, _event=None):
        try:
            raw = self.tree.clipboard_get()
        except tk.TclError:
            return "break"
        text = str(raw or "").strip("\n")
        if not text:
            return "break"

        cols = list(self.tree["columns"])
        for line in text.splitlines():
            parts = line.split("\t")
            if len(parts) == 1 and "," in parts[0]:
                parts = [p.strip() for p in parts[0].split(",")]
            if not parts:
                continue
            vals = (parts + [""] * len(cols))[:len(cols)]
            # status는 시스템 컬럼이므로 강제 PENDING 유지
            try:
                st_idx = cols.index("status")
                vals[st_idx] = "PENDING"
            except ValueError as e:
                logger.warning(f"[_paste_rows_from_clipboard] Suppressed: {e}")
            try:
                q_idx = cols.index("qty_mt")
                q_raw = str(vals[q_idx]).replace(",", "").strip()
                vals[q_idx] = f"{float(q_raw):.4f}" if q_raw else "0.0000"
            except (ValueError, TypeError) as e:
                logger.warning(f"[_paste_rows_from_clipboard] Suppressed: {e}")
            self.tree.insert("", END, values=vals)

        self._after_tree_data_changed()
        return "break"

    def _after_tree_data_changed(self):
        if getattr(self, '_alloc_total_footer', None):
            self._alloc_total_footer.update_totals()
        self.btn_reserve.config(state='normal' if self.tree.get_children("") else 'disabled')

    def _sync_parsed_rows_from_tree(self):
        """트리뷰 현재 데이터를 parsed_rows(dict)로 동기화."""
        rows = []
        for iid in self.tree.get_children(""):
            vals = self.tree.item(iid, "values")
            cols = list(self.tree["columns"])
            d = {cols[i]: (vals[i] if i < len(vals) else "") for i in range(len(cols))}
            lot_no = str(d.get("lot_no", "")).strip()
            if not lot_no:
                continue
            try:
                qty_mt = float(str(d.get("qty_mt", "0")).replace(",", "").strip() or 0.0)
            except (ValueError, TypeError):
                qty_mt = 0.0
            d["qty_mt"] = qty_mt
            d["is_sample"] = (qty_mt > 0 and qty_mt < SAMPLE_MT_THRESHOLD)  # v7.8.0: 샘플 행 복원
            d["sold_to"] = str(d.get("sold_to", "") or d.get("customer", "")).strip()
            rows.append(d)
        self.parsed_rows = rows

    def _select_file(self):
        path = filedialog.askopenfilename(
            parent=self.dialog, title="판매 배정 Excel 선택",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        if path:
            self._file_var.set(path)
            self.source_file = path
            self._parse_file()  # 파일 선택 시 자동 파싱

    def _parse_file(self):
        """Allocation Excel 파싱 — 백그라운드 스레드에서 실행하여 UI 블로킹 방지."""
        path = self._file_var.get()
        if not path:
            CustomMessageBox.showwarning(self.dialog, "경고", "파일을 먼저 선택하세요.")
            return

        self._summary_var.set("⏳ 파일 파싱 중... (잠시 기다려주세요)")
        self.btn_reserve.config(state='disabled')
        self.tree.delete(*self.tree.get_children())

        def _worker():
            result_data = {'result': None, 'elapsed': 0.0, 'error': None}
            try:
                from parsers.allocation_parser import AllocationParser
                t0 = time.perf_counter()
                parser = AllocationParser()
                result_data['result'] = parser.parse(path)
                result_data['elapsed'] = time.perf_counter() - t0
            except (ValueError, TypeError, AttributeError, ImportError) as e:
                result_data['error'] = e
                logger.error(f"Allocation 파싱 오류: {e}", exc_info=True)
            try:
                self.dialog.after(0, lambda: self._apply_parse_result(path, result_data))
            except (tk.TclError, RuntimeError):
                logger.debug("다이얼로그 종료로 apply 스킵")

        threading.Thread(target=_worker, daemon=True).start()

    def _apply_parse_result(self, path: str, data: dict):
        """파싱 결과를 메인 스레드에서 UI에 반영."""
        try:
            if not self.dialog.winfo_exists():
                return
        except (tk.TclError, RuntimeError):
            return
        if data.get('error'):
            self._summary_var.set("파일을 선택하세요")
            CustomMessageBox.showerror(
                self.dialog, "파싱 오류",
                f"Allocation 파일 파싱 실패:\n{data['error']}"
            )
            return

        result = data.get('result')
        elapsed_sec = data.get('elapsed', 0)

        self.parsed_rows = result.rows if result else []

        # v7.8.0: 파싱 결과 0건 + 경고 있음 → 파일 양식 안내
        if not self.parsed_rows and result and getattr(result, 'warnings', []):
            _warn_list = result.warnings
            _warn_msg = "\n".join(str(w) for w in _warn_list[:5])
            CustomMessageBox.showwarning(
                self.dialog, "파일 양식 확인",
                f"Allocation 데이터를 읽을 수 없습니다.\n\n"
                f"이 파일은 Allocation 전용 엑셀이 아닐 수 있습니다.\n"
                f"generated_allocation/ 폴더의 파일을 사용해주세요.\n\n"
                f"상세:\n{_warn_msg}"
            )
            self._summary_var.set("⚠️ 파일 양식 불일치 — Allocation 전용 엑셀을 선택하세요")
            return

        self.tree.delete(*self.tree.get_children())

        for i, row in enumerate(self.parsed_rows):
            vals = (
                getattr(row, 'lot_no', ''),
                getattr(row, 'sap_no', ''),
                getattr(row, 'product', ''),
                f"{getattr(row, 'qty_mt', 0):.4f}",
                getattr(row, 'sold_to', ''),
                getattr(row, 'sale_ref', ''),
                str(getattr(row, 'outbound_date', '') or ''),
                getattr(row, 'warehouse', ''),
                'PENDING',
            )
            self.tree.insert('', END, values=vals)

        header = result.header if result else None
        customer = getattr(header, 'customer', '?') if header else '?'
        total = getattr(result, 'total_qty', None)
        if total is None and self.parsed_rows:
            total = sum(float(getattr(r, 'qty_mt', 0) or 0) for r in self.parsed_rows)
        elif total is None:
            total = 0.0
        fname = path.split('/')[-1].split(chr(92))[-1]
        tb500, samp = _allocation_tonbag_sample_counts(self.parsed_rows)
        self._summary_var.set(
            f"고객: {customer} | 총 {len(self.parsed_rows)}행 | 총량: {total:.4f} MT | "
            f"500kg {tb500}개, 샘플 {samp}개 | 파싱: {elapsed_sec:.2f}초 | {fname}"
        )

        if self.parsed_rows:
            self.btn_reserve.config(state='normal')
            if hasattr(self, 'btn_reset_lot'):
                self.btn_reset_lot.config(state='normal')
        else:
            self.btn_reserve.config(state='disabled')
            if hasattr(self, 'btn_reset_lot'):
                self.btn_reset_lot.config(state='disabled')
        if getattr(self, '_alloc_total_footer', None):
            self._alloc_total_footer.update_totals()

    def _fill_tree_from_parsed_rows(self):
        """parsed_rows( dict 리스트 )로 트리 채우기. show_with_data용."""
        self.tree.delete(*self.tree.get_children())
        for i, row in enumerate(self.parsed_rows):
            if hasattr(row, 'get'):
                vals = (
                    str(row.get('lot_no', '')),
                    str(row.get('sap_no', '')),
                    str(row.get('product', '')),
                    f"{float(row.get('qty_mt') or 0):.4f}",
                    str(row.get('sold_to') or row.get('customer', '')),
                    str(row.get('sale_ref', '')),
                    str(row.get('outbound_date', '') or ''),
                    str(row.get('warehouse', '')),
                    'PENDING',
                )
            else:
                vals = (
                    getattr(row, 'lot_no', ''),
                    getattr(row, 'sap_no', ''),
                    getattr(row, 'product', ''),
                    f"{getattr(row, 'qty_mt', 0):.4f}",
                    getattr(row, 'sold_to', ''),
                    getattr(row, 'sale_ref', ''),
                    str(getattr(row, 'outbound_date', '') or ''),
                    getattr(row, 'warehouse', ''),
                    'PENDING',
                )
            self.tree.insert('', END, values=vals)
        if getattr(self, '_alloc_total_footer', None):
            self._alloc_total_footer.update_totals()

    def _on_reserve(self):
        self._sync_parsed_rows_from_tree()
        if not self.parsed_rows:
            CustomMessageBox.showwarning(self.dialog, "경고", "예약할 행이 없습니다.")
            return

        # 업로드 원본 내부 중복 LOT는 정합성 보호를 위해 전량 제외
        dedup = self._exclude_duplicate_lots_for_integrity(self.parsed_rows)
        excluded_rows = dedup.get('excluded_rows', 0)
        excluded_lots = dedup.get('excluded_lots', [])
        if excluded_rows > 0:
            self.parsed_rows = dedup.get('filtered_rows', [])
            self._fill_tree_from_parsed_rows()
            preview = ", ".join(excluded_lots[:15])
            more = f"\n... 외 {len(excluded_lots) - 15}건" if len(excluded_lots) > 15 else ""
            CustomMessageBox.showwarning(
                self.dialog,
                "중복 LOT 자동 제외",
                "업로드 데이터 내 중복 LOT가 감지되어 해당 화물은 배정에서 제외했습니다.\n\n"
                f"제외 LOT: {len(excluded_lots)}건\n"
                f"제외 행: {excluded_rows}행\n\n"
                f"{preview}{more}"
            )
            if not self.parsed_rows:
                CustomMessageBox.showwarning(self.dialog, "예약 중단", "중복 제외 후 남은 배정 대상이 없습니다.")
                return

        dup = self._check_duplicate_allocation_file()
        if dup.get('is_duplicate'):
            # ⑤ v6.7.1: 중복 파일 안내 강화
            dup_msg = (
                f"⚠️ 동일 파일 중복 업로드 감지\n\n"
                f"파일명: {dup.get('file_name', '(알 수 없음)')}\n"
                f"기존 예약 건수: {dup.get('count', 0)}건\n\n"
                "─────────────────────────────\n"
                "[ 선택 안내 ]\n"
                "  ✅ 예 → 이번 예약을 계속 진행\n"
                "        (기존 예약 + 신규 예약 중복 가능)\n\n"
                "  ❌ 아니오 → 취소\n"
                "        [Allocation 탭] → 기존 예약 선택 취소 후\n"
                "        재업로드 권장\n"
                "─────────────────────────────\n"
                "계속 진행하시겠습니까?"
            )
            if not CustomMessageBox.askyesno(self.dialog, "⚠️ 중복 Allocation 경고", dup_msg):
                return

        warnings = self._build_reserve_shortage_warnings()
        if warnings:
            warn_msg = (
                "[주의: 요청 수량이 현재 판매가능 수량보다 많을 수 있습니다]\n"
                + "\n".join(warnings[:10])
                + "\n\n이 상태에서 진행하면 예약이 실패(롤백)될 수 있습니다.\n그래도 계속하시겠습니까?"
            )
            if not CustomMessageBox.askyesno(self.dialog, "가용 수량 부족", warn_msg):
                return

        history_lots = self._check_reserved_history_lots()
        if history_lots:
            preview = ", ".join(history_lots[:10])
            more = f"\n... 외 {len(history_lots) - 10}건" if len(history_lots) > 10 else ""
            history_msg = (
                "[안내: 이전에 같은 LOT으로 예약 작업을 한 적이 있습니다]\n"
                "아래 LOT은 과거 기록이 있어 중복 작업 가능성을 확인해 주세요.\n"
                f"{preview}{more}\n\n"
                "※ 이 메시지는 과거 기록 안내이며, 현재 재고 부족 오류를 뜻하지는 않습니다.\n\n"
                "이번 예약을 계속 진행할까요?"
            )
            if not CustomMessageBox.askyesno(self.dialog, "예약 이력 경고", history_msg):
                return

        tb500, samp = _allocation_tonbag_sample_counts(self.parsed_rows)
        if self._lot_mode_var.get():
            confirm_msg = (
                f"LOT 단위로 500kg 제품 {tb500}개 및 샘플(1kg) {samp}개를 예약 계획으로 저장합니다.\n"
                "톤백 ID는 지금 지정하지 않으며, 바코드 스캔 시점에 확정됩니다.\n계속하시겠습니까?"
            )
        else:
            confirm_msg = (
                f"500kg 제품 {tb500}개 및 샘플(1kg) {samp}개 판매 배정합니다.\n계속하시겠습니까?"
            )
        ok = CustomMessageBox.askyesno(self.dialog, "예약 실행", confirm_msg)
        if not ok:
            return

        try:
            if hasattr(self.app, 'do_action_tx'):
                result = self.app.do_action_tx(
                    "RESERVE_FROM_ALLOCATION",
                    lambda: self.engine.reserve_from_allocation(
                        self.parsed_rows,
                        source_file=self.source_file,
                        reservation_mode=("lot" if self._lot_mode_var.get() else "tonbag"),
                    ),
                    parent=self.dialog,
                    refresh_mode="deferred",
                )
            else:
                result = self.engine.reserve_from_allocation(
                    self.parsed_rows,
                    source_file=self.source_file,
                    reservation_mode=("lot" if self._lot_mode_var.get() else "tonbag"),
                )
            if result.get('success'):
                reserved = result.get('reserved', 0)
                pending_approval = int(result.get('pending_approval', 0) or 0)
                requested_rows = result.get('requested_rows', len(self.parsed_rows))
                requested_slots = tb500 + samp  # 요청한 톤백/샘플 수 (행 수와 단위 다름)
                msg = (
                    f"✅ 요청: {requested_rows}행 (500kg {tb500}개 + 샘플 {samp}개)\n"
                    f"• 즉시 예약(RESERVED): {reserved}개\n"
                    f"• 승인대기(STAGED): {pending_approval}개"
                )
                is_lot_mode = str(result.get('reservation_mode', '')).strip().lower() == 'lot'
                if pending_approval > 0:
                    msg += (
                        "\n\n(안내) 이번 건은 수량이 커서 바로 예약되지 않고 "
                        "먼저 '승인 대기'로 저장되었습니다."
                    )
                elif requested_slots > 0 and reserved < requested_slots:
                    msg += (
                        f"\n\n(안내) 요청 수량({requested_slots}개)보다 실제 예약이 적습니다.\n"
                        "해당 LOT의 판매가능 톤백이 부족하거나, 이미 예약/출고된 톤백이 포함된 경우입니다."
                    )
                errors = result.get('errors', [])
                if errors:
                    msg += "\n\n(미예약 사유: " + "; ".join(errors[:5]) + ")"
                fail_report = result.get('fail_report') or {}
                if fail_report.get('csv') or fail_report.get('json'):
                    msg += (
                        "\n\n검증 실패 리포트 저장:\n"
                        f"- CSV: {fail_report.get('csv', '-')}\n"
                        f"- JSON: {fail_report.get('json', '-')}"
                    )
                if is_lot_mode:
                    msg += "\n\n[LOT 단위 예약 모드] 톤백 ID는 스캔 시점에 확정됩니다. 출고 메뉴의 바코드 스캔 업로드를 사용하세요."
                if pending_approval > 0:
                    msg += (
                        f"\n\n승인 대기 저장: {pending_approval}건\n"
                        "다음 순서:\n"
                        "1) 출고 > ✅ Allocation 승인 대기\n"
                        "2) 승인\n"
                        "3) 📌 예약 반영 실행"
                    )
                CustomMessageBox.showinfo(self.dialog, "예약 완료", msg)
                self.btn_cancel_res.config(state='normal')
                if is_lot_mode:
                    self.btn_execute.config(state='disabled')
                else:
                    self.btn_execute.config(state='normal')
                if not hasattr(self.app, 'do_action_tx'):
                    self._deferred_refresh_after_action()
                if pending_approval > 0 and hasattr(self.app, '_show_allocation_approval_queue'):
                    # 승인대기 건이 생기면 승인 화면으로 바로 이동
                    try:
                        self.app._show_allocation_approval_queue()
                    except Exception as e:
                        logger.debug(f'Suppressed: {e}')
            else:
                errors = result.get('errors', [])
                fail_report = result.get('fail_report') or {}
                report_msg = ""
                if fail_report.get('csv') or fail_report.get('json'):
                    report_msg = (
                        "\n\n검증 실패 리포트 저장:\n"
                        f"- CSV: {fail_report.get('csv', '-')}\n"
                        f"- JSON: {fail_report.get('json', '-')}"
                    )
                CustomMessageBox.showerror(
                    self.dialog, "예약 실패",
                    "예약 실패:\n" + "\n".join(errors[:10]) + report_msg
                )
                self._show_lot_status_popup()
        except (ValueError, TypeError, AttributeError) as e:
            logger.error(f"예약 실행 오류: {e}", exc_info=True)
            CustomMessageBox.showerror(self.dialog, "오류", f"예약 실행 중 오류:\n{e}")

    def _exclude_duplicate_lots_for_integrity(self, rows: list) -> dict:
        """업로드 데이터 내부 LOT 중복을 모두 제외한다(정합성 우선 정책).

        v7.8.0: 샘플 행(is_sample=True)은 중복 판정에서 제외.
        본 제품 + 샘플이 같은 LOT로 2행 있는 것은 정상 구조.
        """
        if not rows:
            return {"filtered_rows": [], "excluded_rows": 0, "excluded_lots": []}

        # 샘플이 아닌 행의 LOT만으로 중복 판정
        non_sample_lots = []
        for r in rows:
            is_sample = (r.get('is_sample') if hasattr(r, 'get') else getattr(r, 'is_sample', False))
            if is_sample:
                continue
            raw = (r.get('lot_no') if hasattr(r, 'get') else getattr(r, 'lot_no', ''))
            lot = str(raw or '').strip().upper()
            if lot:
                non_sample_lots.append(lot)

        counter = Counter(non_sample_lots)
        duplicate_lot_set = {lot for lot, cnt in counter.items() if cnt > 1}
        if not duplicate_lot_set:
            return {"filtered_rows": rows, "excluded_rows": 0, "excluded_lots": []}

        filtered_rows = []
        excluded_rows = 0
        for r in rows:
            raw = (r.get('lot_no') if hasattr(r, 'get') else getattr(r, 'lot_no', ''))
            lot = str(raw or '').strip().upper()
            if lot and lot in duplicate_lot_set:
                excluded_rows += 1
                continue
            filtered_rows.append(r)

        excluded_lots = sorted(list(duplicate_lot_set))
        logger.warning(
            "[allocation] duplicate lots excluded: lots=%d rows=%d",
            len(excluded_lots),
            excluded_rows,
        )
        return {
            "filtered_rows": filtered_rows,
            "excluded_rows": excluded_rows,
            "excluded_lots": excluded_lots,
        }

    def _check_duplicate_allocation_file(self) -> dict:
        """같은 Allocation 입력이 이미 예약됐는지 확인 (fingerprint 우선)."""
        if not hasattr(self.engine, 'db') or not self.engine.db:
            return {'is_duplicate': False}
        try:
            import os
            source_fp = ""
            if hasattr(self.engine, '_compute_allocation_source_fingerprint'):
                source_fp = self.engine._compute_allocation_source_fingerprint(self.parsed_rows, self.source_file)

            has_fp_col = False
            try:
                info_rows = self.engine.db.fetchall("PRAGMA table_info(allocation_plan)")
                cols = {str(r.get('name', '')).strip().lower() for r in (info_rows or [])}
                has_fp_col = 'source_fingerprint' in cols
            except Exception:
                has_fp_col = False

            if source_fp and has_fp_col:
                row = self.engine.db.fetchone(
                    "SELECT COUNT(*) AS cnt FROM allocation_plan "
                    "WHERE status = 'RESERVED' AND source_fingerprint = ?",
                    (source_fp,)
                )
                cnt = row.get('cnt', 0) if isinstance(row, dict) else (row[0] if row else 0)
                _name = os.path.basename(self.source_file) if self.source_file and self.source_file != '(붙여넣기)' else '(붙여넣기)'
                return {
                    'is_duplicate': cnt > 0,
                    'count': int(cnt),
                    'file_name': _name,
                    'fingerprint': source_fp,
                }

            if not self.source_file or self.source_file == '(붙여넣기)':
                return {'is_duplicate': False}

            fname = os.path.basename(self.source_file)
            row = self.engine.db.fetchone(
                "SELECT COUNT(*) AS cnt FROM allocation_plan WHERE status = 'RESERVED' AND source_file LIKE ?",
                (f"%{fname}",)
            )
            cnt = row.get('cnt', 0) if isinstance(row, dict) else (row[0] if row else 0)
            return {'is_duplicate': cnt > 0, 'count': int(cnt), 'file_name': fname, 'fingerprint': source_fp}
        except Exception as e:
            logger.debug(f"중복 Allocation 확인 실패: {e}")
            return {'is_duplicate': False}

    def _check_reserved_history_lots(self) -> list:
        """과거 예약 이력이 있는 LOT 목록 반환 (allocation_plan 기준)."""
        if not getattr(self, 'parsed_rows', None):
            return []
        if not hasattr(self.engine, 'db') or not self.engine.db:
            return []
        lot_set = []
        for r in self.parsed_rows:
            lot_no = (r.get('lot_no') if hasattr(r, 'get') else getattr(r, 'lot_no', '')).strip()
            if lot_no:
                lot_set.append(lot_no)
        lot_set = list(dict.fromkeys(lot_set))
        if not lot_set:
            return []
        try:
            placeholders = ",".join("?" * len(lot_set))
            rows = self.engine.db.fetchall(
                f"SELECT DISTINCT lot_no FROM allocation_plan WHERE lot_no IN ({placeholders})",
                tuple(lot_set),
            )
            lots = [str(r.get('lot_no', '')).strip() for r in (rows or []) if (r.get('lot_no') if isinstance(r, dict) else None)]
            return [l for l in lots if l]
        except Exception as e:
            logger.debug(f"예약 이력 확인 실패: {e}")
            return []

    def _build_reserve_shortage_warnings(self) -> list:
        """예약 전 LOT별 가용(샘플 포함) 수량 부족 경고 메시지 생성."""
        warnings = []
        if not hasattr(self.engine, 'db') or not self.engine.db:
            return warnings

        by_lot = {}
        for r in self.parsed_rows:
            lot_no = (r.get('lot_no') if hasattr(r, 'get') else getattr(r, 'lot_no', '')).strip()
            if not lot_no:
                continue
            qty_mt = 0.0
            sublot_count = 0
            if hasattr(r, 'get'):
                qty_mt = float(r.get('qty_mt') or 0)
                sublot_count = int(r.get('sublot_count') or r.get('tonbag_count') or 0)
            else:
                qty_mt = float(getattr(r, 'qty_mt', 0) or 0)
                sublot_count = int(getattr(r, 'sublot_count', 0) or getattr(r, 'tonbag_count', 0) or 0)
            weight_kg = qty_mt * 1000.0 if qty_mt > 0 else sublot_count * 500
            need_count = sublot_count if sublot_count > 0 else max(1, int(weight_kg / 500))
            by_lot[lot_no] = by_lot.get(lot_no, 0) + need_count

        for lot_no, need_count in by_lot.items():
            row = self.engine.db.fetchone(
                """SELECT
                    SUM(CASE WHEN status = 'AVAILABLE' THEN 1 ELSE 0 END) AS avail_total,
                    SUM(CASE WHEN status = 'AVAILABLE' AND COALESCE(is_sample, 0) = 1 THEN 1 ELSE 0 END) AS avail_sample
                FROM inventory_tonbag WHERE lot_no = ?""",
                (lot_no,)
            )
            avail_total = row.get('avail_total', 0) if isinstance(row, dict) else (row[0] if row else 0)
            avail_sample = row.get('avail_sample', 0) if isinstance(row, dict) else (row[1] if row else 0)
            if avail_total < need_count:
                warnings.append(
                    f"{lot_no}: 요청 {need_count}개 / 가용(전체) {avail_total}개 "
                    f"(샘플 {avail_sample}개 포함)"
                )

        return warnings

    def _show_lot_status_popup(self) -> None:
        """예약 실패 시 LOT별 상태(샘플 포함) 팝업 표시."""
        if not hasattr(self.engine, 'db') or not self.engine.db:
            return
        lot_nos = []
        for r in self.parsed_rows:
            lot_no = (r.get('lot_no') if hasattr(r, 'get') else getattr(r, 'lot_no', '')).strip()
            if lot_no:
                lot_nos.append(lot_no)
        if not lot_nos:
            return
        # 중복 제거 유지 순서
        seen = set()
        lot_nos = [l for l in lot_nos if not (l in seen or seen.add(l))]

        lines = ["[LOT 상태 요약]", "LOT\tAVAIL\tSAMPLE\tRESERVED\tPICKED\tSOLD"]
        for lot_no in lot_nos:
            row = self.engine.db.fetchone(
                """SELECT
                    SUM(CASE WHEN status = 'AVAILABLE' THEN 1 ELSE 0 END) AS avail_total,
                    SUM(CASE WHEN status = 'AVAILABLE' AND COALESCE(is_sample, 0) = 1 THEN 1 ELSE 0 END) AS avail_sample,
                    SUM(CASE WHEN status = 'RESERVED' THEN 1 ELSE 0 END) AS reserved_cnt,
                    SUM(CASE WHEN status = 'PICKED' THEN 1 ELSE 0 END) AS picked_cnt,
                    SUM(CASE WHEN status = 'SOLD' THEN 1 ELSE 0 END) AS sold_cnt
                FROM inventory_tonbag WHERE lot_no = ?""",
                (lot_no,)
            )
            if not row:
                lines.append(f"{lot_no}: 상태 없음")
                continue
            avail_total = row.get('avail_total', 0)
            avail_sample = row.get('avail_sample', 0)
            reserved_cnt = row.get('reserved_cnt', 0)
            picked_cnt = row.get('picked_cnt', 0)
            sold_cnt = row.get('sold_cnt', 0)
            lines.append(
                f"{lot_no}\t{avail_total}\t{avail_sample}\t{reserved_cnt}\t{picked_cnt}\t{sold_cnt}"
            )
            if len(lines) >= 16:
                lines.append("... (이하 생략)")
                break

        CustomMessageBox.showinfo(self.dialog, "LOT 상태 확인", "\n".join(lines))

    def _on_execute(self):
        """RESERVED → PICKED 전환 (v6.9.5: LOT 모드에서는 스캔 안내)"""
        # v6.9.5 [LOT-MODE]: LOT 단위 예약(tonbag_id=NULL) 상태 확인
        # LOT 모드에서는 바코드 스캔이 유일한 출고 확정 경로
        try:
            lot_mode_cnt = 0
            row = self.engine.db.fetchone(
                "SELECT COUNT(*) AS cnt FROM allocation_plan "
                "WHERE status='RESERVED' AND tonbag_id IS NULL"
            )
            lot_mode_cnt = int(row.get('cnt', 0) if isinstance(row, dict) else (row[0] if row else 0))
        except Exception:
            lot_mode_cnt = 0

        if lot_mode_cnt > 0:
            CustomMessageBox.showinfo(
                self.dialog, "📡 바코드 스캔 필요",
                f"LOT 단위 예약 {lot_mode_cnt}건이 바코드 스캔 대기 중입니다.\n\n"
                f"▶ 출고 실행은 바코드 스캔으로만 확정됩니다.\n"
                f"   메뉴 → 출고 스캔 → 톤백 바코드를 스캔하세요.\n\n"
                f"(이 버튼은 tonbag 직접 지정 방식에서만 사용합니다)"
            )
            return

        ok = CustomMessageBox.askyesno(
            self.dialog, "출고 실행",
            "RESERVED 상태의 톤백을 PICKED로 전환합니다.\n계속하시겠습니까?"
        )
        if not ok:
            return
        try:
            if hasattr(self.app, 'do_action_tx'):
                result = self.app.do_action_tx(
                    "EXECUTE_RESERVED",
                    self.engine.execute_reserved,
                    parent=self.dialog,
                    refresh_mode="deferred",
                )
            else:
                result = self.engine.execute_reserved()
            if result.get('success'):
                CustomMessageBox.showinfo(
                    self.dialog, "출고 실행 완료",
                    f"✅ {result.get('executed', 0)}개 톤백 PICKED 전환 완료"
                )
                self.btn_confirm.config(state='normal')
                if not hasattr(self.app, 'do_action_tx'):
                    self._deferred_refresh_after_action()
            else:
                CustomMessageBox.showerror(
                    self.dialog, "출고 실행 실패",
                    "\n".join(result.get('errors', ['알 수 없는 오류']))
                )
        except (ValueError, TypeError, AttributeError) as e:
            logger.error(f"출고 실행 오류: {e}", exc_info=True)
            CustomMessageBox.showerror(self.dialog, "오류", str(e))

    def _on_confirm(self):
        """PICKED → SOLD 확정 (v6.9.5: LOT 모드 잔여 예약 경고)"""
        # v6.9.5 [LOT-MODE]: 바코드 스캔 대기 중인 예약이 있으면 경고
        try:
            pending_cnt = 0
            row = self.engine.db.fetchone(
                "SELECT COUNT(*) AS cnt FROM allocation_plan "
                "WHERE status='RESERVED' AND tonbag_id IS NULL"
            )
            pending_cnt = int(row.get('cnt', 0) if isinstance(row, dict) else (row[0] if row else 0))
            if pending_cnt > 0:
                proceed = CustomMessageBox.askyesno(
                    self.dialog, "⚠️ 스캔 미완료 예약 있음",
                    f"바코드 스캔 대기 중인 예약 {pending_cnt}건이 있습니다.\n\n"
                    f"스캔 없이 확정하면 해당 톤백은 SOLD 처리되지 않습니다.\n"
                    f"그래도 진행하시겠습니까?"
                )
                if not proceed:
                    return
        except Exception:
            pass

        ok = CustomMessageBox.askyesno(
            self.dialog, "출고 확정",
            "PICKED 상태의 톤백을 SOLD로 확정합니다.\n이 작업은 되돌릴 수 없습니다.\n계속하시겠습니까?"
        )
        if not ok:
            return
        try:
            if hasattr(self.app, 'do_action_tx'):
                result = self.app.do_action_tx(
                    "CONFIRM_OUTBOUND",
                    self.engine.confirm_outbound,
                    parent=self.dialog,
                    refresh_mode="deferred",
                )
            else:
                result = self.engine.confirm_outbound()
            if result.get('success'):
                CustomMessageBox.showinfo(
                    self.dialog, "출고 확정 완료",
                    f"✅ {result.get('confirmed', 0)}개 톤백 SOLD 확정"
                )
                if not hasattr(self.app, 'do_action_tx'):
                    self._deferred_refresh_after_action()
            else:
                CustomMessageBox.showerror(
                    self.dialog, "확정 실패",
                    "\n".join(result.get('errors', ['알 수 없는 오류']))
                )
        except (ValueError, TypeError, AttributeError) as e:
            logger.error(f"출고 확정 오류: {e}", exc_info=True)
            CustomMessageBox.showerror(self.dialog, "오류", str(e))

    def _on_cancel_reservation(self):
        """④ v6.7.1: 예약 취소 — 전체 취소 전 이중 경고 추가"""
        ok = CustomMessageBox.askyesno(
            self.dialog, "⚠️ 전체 예약 취소 경고",
            "RESERVED 상태의 톤백을 전체 AVAILABLE로 되돌립니다.\n\n"
            "💡 특정 건만 취소하려면 [Allocation 탭]에서\n"
            "   해당 LOT 우클릭 → 선택 취소를 사용하세요.\n\n"
            "전체 취소를 계속하시겠습니까?"
        )
        if not ok:
            return
        try:
            if hasattr(self.app, 'do_action_tx'):
                result = self.app.do_action_tx(
                    "CANCEL_RESERVATION",
                    self.engine.cancel_reservation,
                    parent=self.dialog,
                    refresh_mode="deferred",
                )
            else:
                result = self.engine.cancel_reservation()
            if result.get('success'):
                CustomMessageBox.showinfo(
                    self.dialog, "예약 취소 완료",
                    f"✅ {result.get('cancelled', 0)}개 톤백 예약 취소됨"
                )
                if not hasattr(self.app, 'do_action_tx'):
                    self._deferred_refresh_after_action()
            else:
                CustomMessageBox.showerror(
                    self.dialog, "취소 실패",
                    "\n".join(result.get('errors', ['알 수 없는 오류']))
                )
        except (ValueError, TypeError, AttributeError) as e:
            logger.error(f"예약 취소 오류: {e}", exc_info=True)
            CustomMessageBox.showerror(self.dialog, "오류", str(e))

    def _on_reset_reservation_for_lots(self) -> None:
        """현재 Allocation LOT들에 대한 예약(RESERVED)만 초기화."""
        if not self.parsed_rows:
            return
        lot_nos = []
        for r in self.parsed_rows:
            lot_no = (r.get('lot_no') if hasattr(r, 'get') else getattr(r, 'lot_no', '')).strip()
            if lot_no:
                lot_nos.append(lot_no)
        if not lot_nos:
            return
        # 중복 제거 유지 순서
        seen = set()
        lot_nos = [l for l in lot_nos if not (l in seen or seen.add(l))]

        if not CustomMessageBox.askyesno(
            self.dialog, "LOT 예약 초기화",
            f"현재 Allocation LOT {len(lot_nos)}개에 대해\n"
            f"RESERVED 상태만 AVAILABLE로 되돌립니다.\n"
            f"(PICKED/SOLD는 변경하지 않습니다)\n\n계속하시겠습니까?"
        ):
            return
        total = 0
        try:
            for lot_no in lot_nos:
                r = self.engine.cancel_reservation(lot_no=lot_no)
                total += r.get('cancelled', 0)
            CustomMessageBox.showinfo(
                self.dialog, "초기화 완료",
                f"예약 초기화 완료: {total}건 (RESERVED → AVAILABLE)"
            )
            if hasattr(self.app, 'refresh_bus_deferred'):
                self.app.refresh_bus_deferred(reason="RESET_RESERVATION_FOR_LOTS", delay_ms=50)
            else:
                self._deferred_refresh_after_action()
        except (ValueError, TypeError, AttributeError) as e:
            logger.error(f"LOT 예약 초기화 오류: {e}", exc_info=True)
            CustomMessageBox.showerror(self.dialog, "오류", str(e))

    def _show_reservation_status(self):
        """allocation_plan 예약 현황 (필터 + 엑셀 저장)."""
        try:
            try:
                info_rows = self.engine.db.fetchall("PRAGMA table_info(allocation_plan)")
                plan_cols = {str(r.get('name', '')).strip().lower() for r in (info_rows or [])}
            except Exception:
                plan_cols = set()
            has_source_col = 'source' in plan_cols

            status_win = tk.Toplevel(self.dialog)
            status_win.title("📊 판매 배정 예약 현황")
            setup_dialog_geometry_persistence(status_win, "allocation_status_win", self.dialog, "large")
            _is_dark = ThemeColors.is_dark_theme(getattr(self.app, 'current_theme', 'flatly'))
            status_win.configure(bg=ThemeColors.get('bg_card', _is_dark))
            status_win.transient(self.dialog)

            today = time.strftime("%Y-%m-%d")
            customer_var = tk.StringVar()
            lot_var = tk.StringVar()
            start_var = tk.StringVar(value=today)
            end_var = tk.StringVar(value=today)
            summary_var = tk.StringVar(value="")
            state = {"rows": []}

            filter_frm = ttk.LabelFrame(status_win, text="필터 (고객사/LOT/기간)")
            filter_frm.pack(fill=X, padx=5, pady=(5, 2))
            ttk.Label(filter_frm, text="고객사").grid(row=0, column=0, padx=4, pady=4, sticky="w")
            # v6.7.3: 고객사 Combobox (광양 거래처 preset)
            _cust_cb = ttk.Combobox(
                filter_frm, textvariable=customer_var, width=22,
                values=[''] + CUSTOMER_PRESETS, state='normal'
            )
            _cust_cb.grid(row=0, column=1, padx=4, pady=4, sticky="w")
            ttk.Label(filter_frm, text="LOT").grid(row=0, column=2, padx=4, pady=4, sticky="w")
            ttk.Entry(filter_frm, textvariable=lot_var, width=16).grid(row=0, column=3, padx=4, pady=4, sticky="w")
            ttk.Label(filter_frm, text="시작일").grid(row=0, column=4, padx=4, pady=4, sticky="w")
            ttk.Entry(filter_frm, textvariable=start_var, width=12).grid(row=0, column=5, padx=4, pady=4, sticky="w")
            ttk.Label(filter_frm, text="종료일").grid(row=0, column=6, padx=4, pady=4, sticky="w")
            ttk.Entry(filter_frm, textvariable=end_var, width=12).grid(row=0, column=7, padx=4, pady=4, sticky="w")

            cols = ('lot_no', 'customer', 'sale_ref', 'qty_mt', 'outbound_date', 'status', 'res_mode', 'plan_count', 'tb_count', 'created_at')
            hdrs = ('LOT NO', 'CUSTOMER', 'SALE REF', 'QTY(MT)', 'DATE', 'STATUS', 'MODE', 'PLANS', 'TONBAGS', 'CREATED')
            tree_frame = ttk.Frame(status_win)
            tree_frame.pack(fill=BOTH, expand=True, padx=5, pady=5)
            st = ttk.Treeview(tree_frame, columns=cols, show='headings', height=15)
            scroll = tk.Scrollbar(tree_frame, orient=VERTICAL, command=st.yview)
            scroll_x = tk.Scrollbar(tree_frame, orient='horizontal', command=st.xview)
            st.configure(yscrollcommand=scroll.set, xscrollcommand=scroll_x.set)
            st.pack(side=LEFT, fill=BOTH, expand=True)
            scroll.pack(side=RIGHT, fill=Y)
            scroll_x.pack(side=tk.BOTTOM, fill=tk.X)
            for c, h in zip(cols, hdrs):
                st.heading(c, text=h)
                st.column(c, width=100, anchor='center')

            footer = ttk.Frame(status_win, padding=(5, 4))
            footer.pack(fill=tk.X)
            ttk.Label(
                footer,
                textvariable=summary_var,
                font=('맑은 고딕', 10, 'bold'),
            ).pack(anchor=tk.W)

            def _fetch_rows():
                sql_where = ["ap.status != 'CANCELLED'"]
                params = []
                customer_kw = customer_var.get().strip()
                if customer_kw:
                    sql_where.append("COALESCE(ap.customer, '') LIKE ?")
                    params.append(f"%{customer_kw}%")
                lot_kw = lot_var.get().strip()
                if lot_kw:
                    sql_where.append("ap.lot_no LIKE ?")
                    params.append(f"%{lot_kw}%")
                start_d = start_var.get().strip()
                end_d = end_var.get().strip()
                if start_d:
                    sql_where.append("date(ap.created_at) >= date(?)")
                    params.append(start_d)
                if end_d:
                    sql_where.append("date(ap.created_at) <= date(?)")
                    params.append(end_d)
                where_clause = " AND ".join(sql_where)

                if has_source_col:
                    query = f"""
                        SELECT ap.lot_no, ap.customer, ap.sale_ref, SUM(COALESCE(ap.qty_mt, 0)) AS qty_mt,
                               ap.outbound_date, ap.status, ap.source_file, ap.source,
                               MAX(ap.created_at) AS created_at,
                               COUNT(*) as plan_count,
                               SUM(CASE WHEN ap.tonbag_id IS NOT NULL THEN 1 ELSE 0 END) as tb_count,
                               SUM(CASE WHEN ap.tonbag_id IS NULL THEN 1 ELSE 0 END) as lot_plan_count
                        FROM allocation_plan ap
                        WHERE {where_clause}
                        GROUP BY ap.lot_no, ap.customer, ap.sale_ref, ap.status, ap.source
                        ORDER BY created_at DESC
                    """
                else:
                    query = f"""
                        SELECT ap.lot_no, ap.customer, ap.sale_ref, SUM(COALESCE(ap.qty_mt, 0)) AS qty_mt,
                               ap.outbound_date, ap.status, ap.source_file,
                               MAX(ap.created_at) AS created_at,
                               COUNT(*) as plan_count,
                               SUM(CASE WHEN ap.tonbag_id IS NOT NULL THEN 1 ELSE 0 END) as tb_count,
                               SUM(CASE WHEN ap.tonbag_id IS NULL THEN 1 ELSE 0 END) as lot_plan_count
                        FROM allocation_plan ap
                        WHERE {where_clause}
                        GROUP BY ap.lot_no, ap.customer, ap.sale_ref, ap.status
                        ORDER BY created_at DESC
                    """
                return self.engine.db.fetchall(query, tuple(params)) or []

            def _render_rows(rows):
                state["rows"] = rows
                st.delete(*st.get_children())
                for r in rows:
                    src = str(r.get('source', '') or '').strip().upper()
                    lot_plans = int(r.get('lot_plan_count', 0) or 0)
                    res_mode = 'LOT' if (src == 'LOT' or lot_plans > 0) else 'TONBAG'
                    st.insert('', END, values=(
                        r.get('lot_no', ''), r.get('customer', ''),
                        r.get('sale_ref', ''), f"{float(r.get('qty_mt', 0) or 0):.4f}",
                        r.get('outbound_date', ''), r.get('status', ''),
                        res_mode,
                        r.get('plan_count', 0),
                        r.get('tb_count', 0),
                        r.get('created_at', '')
                    ))
                total_mt = sum(float(r.get('qty_mt', 0) or 0) for r in rows)
                total_tb = sum(int(r.get('tb_count', 0) or 0) for r in rows)
                total_plans = sum(int(r.get('plan_count', 0) or 0) for r in rows)
                lot_mode_open = sum(
                    int(r.get('lot_plan_count', 0) or 0)
                    for r in rows
                    if str(r.get('status', '')).strip().upper() == STATUS_RESERVED
                )
                summary_var.set(
                    f"건수: {len(rows)} LOT  |  계획: {total_plans}건  |  톤백: {total_tb}개  |  "
                    f"LOT모드 잔량(RESERVED): {lot_mode_open}건  |  QTY(MT) 합계: {total_mt:,.4f}"
                )

            def _on_search():
                rows = _fetch_rows()
                _render_rows(rows)
                if not rows:
                    CustomMessageBox.showinfo(status_win, "조회 결과", "조건에 맞는 예약 항목이 없습니다.")

            def _on_reset():
                customer_var.set("")
                lot_var.set("")
                start_var.set(today)
                end_var.set(today)
                _render_rows(_fetch_rows())

            def _on_export_excel():
                rows = state.get("rows", [])
                if not rows:
                    CustomMessageBox.showwarning(status_win, "저장 불가", "저장할 데이터가 없습니다.")
                    return
                fpath = filedialog.asksaveasfilename(
                    parent=status_win,
                    title="예약 현황 Excel 저장",
                    defaultextension=".xlsx",
                    filetypes=[("Excel files", "*.xlsx")],
                    initialfile=f"allocation_status_{time.strftime('%Y%m%d_%H%M%S')}.xlsx",
                )
                if not fpath:
                    return
                try:
                    from openpyxl import Workbook
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "allocation_status"
                    ws.append(['LOT NO', 'CUSTOMER', 'SALE REF', 'QTY(MT)', 'DATE', 'STATUS', 'MODE', 'PLANS', 'TONBAGS', 'CREATED'])
                    for r in rows:
                        src = str(r.get('source', '') or '').strip().upper()
                        lot_plans = int(r.get('lot_plan_count', 0) or 0)
                        res_mode = 'LOT' if (src == 'LOT' or lot_plans > 0) else 'TONBAG'
                        ws.append([
                            r.get('lot_no', ''),
                            r.get('customer', ''),
                            r.get('sale_ref', ''),
                            float(r.get('qty_mt', 0) or 0),
                            r.get('outbound_date', ''),
                            r.get('status', ''),
                            res_mode,
                            int(r.get('plan_count', 0) or 0),
                            int(r.get('tb_count', 0) or 0),
                            r.get('created_at', ''),
                        ])
                    wb.save(fpath)
                    CustomMessageBox.showinfo(status_win, "저장 완료", f"Excel 저장 완료\n{fpath}")
                except Exception as e:
                    logger.error(f"예약 현황 Excel 저장 오류: {e}", exc_info=True)
                    CustomMessageBox.showerror(status_win, "저장 실패", str(e))

            btn_frm = ttk.Frame(status_win)
            btn_frm.pack(fill=X, padx=5, pady=(0, 6))
            ttk.Button(btn_frm, text="조회", command=_on_search).pack(side=LEFT, padx=2)
            ttk.Button(btn_frm, text="초기화", command=_on_reset).pack(side=LEFT, padx=2)
            ttk.Button(btn_frm, text="Excel 저장", command=_on_export_excel).pack(side=LEFT, padx=2)
            ttk.Button(btn_frm, text="닫기", command=status_win.destroy).pack(side=RIGHT, padx=2)

            _render_rows(_fetch_rows())

        except (ValueError, TypeError, AttributeError) as e:
            logger.error(f"예약 현황 조회 오류: {e}", exc_info=True)
            CustomMessageBox.showerror(self.dialog, "오류", str(e))

    def _deferred_refresh_after_action(self):
        """grab 해제 후 50ms 지연으로 리프레시 — 모달 안에서의 무거운 Treeview 갱신 블로킹 방지."""
        try:
            self.dialog.grab_release()
        except (tk.TclError, RuntimeError) as e:
            logger.warning(f"[_deferred_refresh_after_action] Suppressed: {e}")
        if hasattr(self.app, 'refresh_bus_deferred'):
            self.app.refresh_bus_deferred(reason="ALLOCATION_DIALOG_ACTION", delay_ms=50)
        self.app._safe_refresh()
    def _refresh_after_action(self):
        """예약/출고/취소 후 앱 전체 탭 새로고침 (판매 배정/판매화물 결정/출고 탭 포함)"""
        try:
            if hasattr(self.app, 'refresh_bus'):
                self.app.refresh_bus(reason="ALLOCATION_DIALOG_ACTION_IMMEDIATE")
            elif hasattr(self.app, '_refresh_main_tabs'):
                self.app._refresh_main_tabs()
            else:
                if hasattr(self.app, '_refresh_inventory'):
                    self.app._refresh_inventory()
                if hasattr(self.app, '_refresh_tonbag'):
                    self.app._refresh_tonbag()
                if hasattr(self.app, '_refresh_outbound_scheduled'):
                    self.app._refresh_outbound_scheduled()
                if hasattr(self.app, '_refresh_allocation'):
                    self.app._refresh_allocation()
                if hasattr(self.app, '_refresh_picked'):
                    self.app._refresh_picked()
                if hasattr(self.app, '_refresh_sold'):
                    self.app._refresh_sold()
        except (RuntimeError, ValueError) as e:
            logger.debug(f"새로고침 실패: {e}")
